from __future__ import annotations

import pandas as pd
from io import BytesIO
from openpyxl import load_workbook

from src.eklaim_analyzer import EklaimAnalysisResult
from src.eklaim_exporter import export_eklaim_analysis_to_excel
from src.eklaim_formatting import (
    format_analysis_frame_for_display,
    format_idr,
    format_percentage,
    format_summary_value,
)


def test_format_idr_uses_indonesian_rupiah_format() -> None:
    assert format_idr(1_000_000) == "Rp 1.000.000,00"
    assert format_idr(-2500.5) == "-Rp 2.500,50"


def test_format_percentage_appends_percent_suffix() -> None:
    assert format_percentage(30) == "30,00%"
    assert format_percentage(12.345) == "12,35%"


def test_format_summary_value_formats_tariff_metrics() -> None:
    assert format_summary_value("Total Tarif RS", 2_500_000) == "Rp 2.500.000,00"
    assert format_summary_value("Total Klaim Rawat Inap", 12) == "12"


def test_format_analysis_frame_for_display_formats_tariff_and_percent_columns() -> None:
    frame = pd.DataFrame(
        [
            {
                "TOTAL_TARIF": 1_000_000,
                "TARIF_RS": 1_500_000,
                "Selisih_Rp": 500_000,
                "Selisih_Pct": 33.33,
            }
        ]
    )
    display = format_analysis_frame_for_display(frame)
    assert display.iloc[0]["TOTAL_TARIF"] == "Rp 1.000.000,00"
    assert display.iloc[0]["Selisih_Pct"] == "33,33%"


def test_excel_export_keeps_tariffs_numeric_and_formats_percent_as_percent_points() -> None:
    result = EklaimAnalysisResult(
        summary={"Total Tarif RS": 1_500_000, "Total Klaim Rawat Inap": 1},
        selisih_gt_30pct_df=pd.DataFrame([{"TARIF_RS": 1_500_000, "Selisih_Pct": 33.33}]),
    )
    workbook = load_workbook(BytesIO(export_eklaim_analysis_to_excel(result)))
    summary = workbook["ringkasan"]
    assert summary["B2"].value == 1_500_000
    assert isinstance(summary["B2"].value, int)
    assert summary["B2"].number_format.startswith('"Rp "')
    flagged = workbook["selisih_lebih_30pct"]
    assert flagged["A2"].value == 1_500_000
    assert flagged["B2"].value == 0.3333
    assert flagged["B2"].number_format == "0.00%"


def test_nan_is_blank_in_idr_display_and_excel():
    assert format_idr(float("nan")) == ""
    frame = pd.DataFrame([{"TARIF_RS": 1_500_000.0, "Selisih_Pct": float("nan")}])
    assert format_analysis_frame_for_display(frame).iloc[0]["Selisih_Pct"] == ""
    result = EklaimAnalysisResult(
        warnings=["Tarif parsial"],
        invalid_ptd_df=pd.DataFrame(columns=["SEP", "PTD"]),
        selisih_gt_30pct_df=frame,
    )
    workbook = load_workbook(BytesIO(export_eklaim_analysis_to_excel(result)))
    assert workbook["selisih_lebih_30pct"]["A2"].value == 1_500_000
    assert workbook["selisih_lebih_30pct"]["B2"].value is None
    assert workbook["peringatan"]["A2"].value == "Tarif parsial"
    assert workbook["ptd_tidak_valid"]["A1"].value == "SEP"


def test_metadata_display_uses_text_without_changing_export_numbers():
    frame = pd.DataFrame({"Metrik": ["Sumber", "Baris masuk", "Kosong"], "Nilai": ["uji.txt", 2, None]})
    display = format_analysis_frame_for_display(frame)
    assert display["Nilai"].tolist() == ["uji.txt", "2", ""]
    workbook = load_workbook(BytesIO(export_eklaim_analysis_to_excel(EklaimAnalysisResult(metadata_df=frame))))
    assert workbook["metadata"]["B3"].value == 2
