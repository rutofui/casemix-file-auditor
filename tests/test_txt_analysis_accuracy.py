from __future__ import annotations

import io
import json
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.eklaim_analyzer import build_eklaim_analysis
from src.eklaim_exporter import export_eklaim_analysis_to_excel
from src.parser_eklaim_txt import REQUIRED_COLUMNS, TARIFF_COMPONENTS, read_eklaim_txt


SEP = "0132R0770626V000001"


def _claim(**overrides: str) -> dict[str, str]:
    row = dict.fromkeys(REQUIRED_COLUMNS, "")
    row.update(
        SEP=SEP,
        PTD="1",
        INACBG="K-4-17-I",
        DIAGLIST="A09.9;E86",
        PROCLIST="90.59",
        TOTAL_TARIF="1000",
        TARIF_RS="1200",
        LOS="2",
        DPJP="dr. Uji",
        C2=json.dumps({"idrg": {"cost_weight": 0.5, "total_tarif": 1100}}),
        NAMA_PASIEN="Pasien Uji",
        MRN="RM001",
        NOKARTU="KARTU001",
        ADMISSION_DATE="01/08/2026",
        DISCHARGE_DATE="02/08/2026",
    )
    row.update(dict.fromkeys(TARIFF_COMPONENTS, "0"))
    row.update(overrides)
    return row


def _parsed(*rows: dict[str, str], source: str = "fixture.txt") -> pd.DataFrame:
    columns = list(dict.fromkeys([*REQUIRED_COLUMNS, "TARIF_INACBG", "NOKARTU", "ADMISSION_DATE", "DISCHARGE_DATE", *TARIFF_COMPONENTS]))
    payload = pd.DataFrame(rows, columns=columns).fillna("").to_csv(sep="\t", index=False).encode()
    return read_eklaim_txt(io.BytesIO(payload), source_label=source).df


def test_exact_duplicates_are_counted_once_across_input_files() -> None:
    one = _parsed(_claim(), source="one.txt")
    two = _parsed(_claim(), source="two.txt")

    result = build_eklaim_analysis(one, two)

    assert result.summary["Total Klaim Keseluruhan"] == 1
    assert result.data_quality["Salinan Duplikat Identik"] == 1
    assert len(result.duplicate_claims_df) == 2
    assert len(result.quarantine_df) == 1


def test_conflicting_rows_with_same_sep_are_all_quarantined() -> None:
    result = build_eklaim_analysis(
        _parsed(_claim(TARIF_RS="1200")),
        _parsed(_claim(TARIF_RS="1300")),
    )

    assert result.summary["Total Klaim Keseluruhan"] == 0
    assert result.data_quality["SEP Konflik"] == 1
    assert len(result.quarantine_df) == 2
    assert result.tariff_comparison_df.loc[
        result.tariff_comparison_df["Jenis Rawat"] == "Keseluruhan", "Jumlah Klaim"
    ].item() == 0


def test_invalid_sep_and_ptd_never_enter_analysis_kpis() -> None:
    result = build_eklaim_analysis(
        _parsed(_claim(), _claim(SEP="tidak-valid"), _claim(SEP="0132R0770626V000002", PTD="9")),
        pd.DataFrame(),
    )

    assert result.summary["Total Klaim Keseluruhan"] == 1
    assert result.data_quality["Baris Dikecualikan"] == 2
    assert set(result.quarantine_df["SEP"]) == {"tidak-valid", "0132R0770626V000002"}


def test_tariff_comparison_keeps_both_bases_and_uses_paired_denominators() -> None:
    frame = _parsed(
        _claim(TOTAL_TARIF="1000", C2=json.dumps({"idrg": {"total_tarif": 1100, "cost_weight": 0.5}})),
        _claim(SEP="0132R0770626V000002", TOTAL_TARIF="2000", TARIF_RS="", C2=json.dumps({"idrg": {"total_tarif": 2200}})),
    )
    result = build_eklaim_analysis(frame, pd.DataFrame())
    overall = result.tariff_comparison_df.query("`Jenis Rawat` == 'Keseluruhan'").iloc[0]

    assert overall["Total Tarif INA-CBG"] == 3000
    assert overall["Total Tarif iDRG"] == 3300
    assert overall["Klaim Tarif iDRG"] == 2
    assert overall["Pasangan RS - iDRG"] == 1
    assert overall["Selisih RS - iDRG"] == 100
    assert overall["Pasangan iDRG - INA-CBG"] == 2
    assert overall["Selisih iDRG - INA-CBG"] == 300
    assert result.casemix_index["Rawat Inap"]["Casemix Index"] == pytest.approx(0.5)


def test_total_cost_weight_is_rounded_for_display_and_export() -> None:
    frame = _parsed(
        _claim(C2=json.dumps({"idrg": {"cost_weight": 0.5, "total_cost_weight": 0.1}})),
        _claim(
            SEP="0132R0770626V000002",
            C2=json.dumps({"idrg": {"cost_weight": 0.5, "total_cost_weight": 0.2}}),
        ),
    )
    result = build_eklaim_analysis(frame, pd.DataFrame())

    assert result.casemix_index["Rawat Inap"]["Total total_cost_weight"] == 0.3


def test_missing_cost_weight_does_not_hide_idrg_tariff_and_all_missing_cmi_is_none() -> None:
    frame = _parsed(_claim(C2=json.dumps({"idrg": {"total_tarif": 1100}})))
    result = build_eklaim_analysis(frame, pd.DataFrame())
    overall = result.tariff_comparison_df.query("`Jenis Rawat` == 'Keseluruhan'").iloc[0]

    assert overall["Total Tarif iDRG"] == 1100
    cmi = result.casemix_index["Rawat Inap"]
    assert cmi["Casemix Index"] is None
    assert cmi["Jumlah Klaim"] == 1
    assert cmi["Klaim dengan Cost Weight"] == 0
    assert cmi["Total Cost Weight"] is None


def test_mixed_idrg_versions_suppress_aggregate_cmi_but_keep_version_profiles() -> None:
    frame = _parsed(
        _claim(C2=json.dumps({"idrg": {"cost_weight": 0.4, "total_tarif": 1100, "grouper_version": "v1"}})),
        _claim(SEP="0132R0770626V000002", C2=json.dumps({"idrg": {"cost_weight": 0.6, "total_tarif": 1200, "grouper_version": "v2"}})),
    )
    result = build_eklaim_analysis(frame, pd.DataFrame())
    cmi = result.casemix_index["Rawat Inap"]
    profiles = result.casemix_profile_df.query("Sistem == 'iDRG' and `Jenis Rawat` == 'Rawat Inap'")

    assert cmi["Casemix Index"] is None
    assert cmi["Status CMI"] == "Versi berbeda; lihat profil per versi"
    assert profiles["Versi Grouper"].nunique() == 2
    assert set(profiles["Casemix Index"].dropna()) == {0.4, 0.6}


def test_activity_components_and_code_prevalence_reconcile() -> None:
    frame = _parsed(
        _claim(DIAGLIST="A09.9;A09.9;E86", PROSEDUR_BEDAH="100", LABORATORIUM="50", TARIF_RS="1200"),
        _claim(SEP="0132R0770626V000002", MRN="RM002", NOKARTU="KARTU002", PROSEDUR_BEDAH="300", LABORATORIUM="150", TARIF_RS="1800"),
    )
    result = build_eklaim_analysis(frame, pd.DataFrame())
    activity = result.activity_df.query("`Jenis Rawat` == 'Rawat Inap'").iloc[0]
    surgery = result.components_df.query("Komponen == 'PROSEDUR_BEDAH' and `Jenis Rawat` == 'Rawat Inap'").iloc[0]
    diagnosis = result.top_icd10_ri_df.set_index("Kode")

    assert activity["Jumlah Klaim"] == 2
    assert activity["Pasien Unik"] == 2
    assert activity["Total LOS"] == 4
    assert activity["ALOS"] == pytest.approx(2)
    assert surgery["Total Tarif Komponen"] == 400
    assert surgery["Klaim dengan Komponen"] == 2
    assert surgery["Porsi Tarif RS (%)"] == pytest.approx(400 / 3000 * 100)
    assert surgery["Tarif Komponen per Klaim"] == pytest.approx(200)
    assert surgery["Tarif Komponen per Hari"] == pytest.approx(100)
    assert pd.api.types.is_numeric_dtype(result.components_df["Total Tarif Komponen"])
    assert diagnosis.loc["A09.9", "Frekuensi"] == 3
    assert diagnosis.loc["A09.9", "Jumlah Klaim"] == 2
    assert diagnosis.loc["A09.9", "Prevalensi (%)"] == pytest.approx(100)


def test_period_activity_does_not_claim_mom_change_across_missing_month() -> None:
    frame = _parsed(
        _claim(ADMISSION_DATE="30/01/2026", DISCHARGE_DATE="31/01/2026"),
        _claim(SEP="0132R0770626V000002", ADMISSION_DATE="01/03/2026", DISCHARGE_DATE="01/03/2026"),
        _claim(SEP="0132R0770626V000003", ADMISSION_DATE="02/03/2026", DISCHARGE_DATE="02/03/2026"),
    )
    result = build_eklaim_analysis(frame, pd.DataFrame())
    periods = result.period_activity_df.query("`Jenis Rawat` == 'Rawat Inap'").set_index("Periode")

    assert list(periods.index) == ["2026-01", "2026-03"]
    assert pd.isna(periods.loc["2026-01", "Perubahan Klaim (%)"])
    assert pd.isna(periods.loc["2026-03", "Perubahan Klaim (%)"])
    assert periods.loc["2026-03", "Keterangan Periode"] == "Tidak ada bulan sebelumnya yang berurutan"


def test_dates_los_and_blank_patient_id_are_reported_without_los_correction() -> None:
    frame = _parsed(_claim(MRN="", NOKARTU="", LOS="5"))
    result = build_eklaim_analysis(frame, pd.DataFrame())
    issues = result.validation_df

    assert set(issues["Field"]) >= {"MRN", "LOS"}
    assert issues.loc[issues["Field"] == "LOS", "Masalah"].str.contains("tanggal", case=False).any()
    assert result.activity_df.query("`Jenis Rawat` == 'Rawat Inap'")["Total LOS"].item() == 5


def test_reversed_dates_are_excluded_from_period_and_daily_profiles() -> None:
    frame = _parsed(_claim(ADMISSION_DATE="03/01/2026", DISCHARGE_DATE="01/01/2026"))
    result = build_eklaim_analysis(frame, pd.DataFrame())

    assert result.period_activity_df.empty
    assert result.service_days_df.empty
    assert result.validation_df["Masalah"].str.contains("mendahului", case=False).any()


def test_source_row_numbers_are_preserved_in_quarantine_and_duplicate_details() -> None:
    result = build_eklaim_analysis(
        _parsed(_claim(), _claim(TARIF_RS="1300"), source="klaim.txt"),
        pd.DataFrame(),
    )

    assert set(result.quarantine_df["Sumber"]) == {"klaim.txt"}
    assert set(result.quarantine_df["Baris"]) == {2, 3}
    assert set(result.duplicate_claims_df["Baris"]) == {2, 3}


def test_all_quarantined_input_returns_zero_analysis_without_crashing() -> None:
    result = build_eklaim_analysis(_parsed(_claim(SEP="tidak-valid")), pd.DataFrame())

    assert result.summary["Total Klaim Keseluruhan"] == 0
    assert result.data_quality["Klaim Dianalisis"] == 0
    assert len(result.quarantine_df) == 1
    assert result.tariff_comparison_df["Jumlah Klaim"].sum() == 0


def test_invalid_numeric_domain_is_reported_and_omitted_from_tariff_totals() -> None:
    result = build_eklaim_analysis(
        _parsed(_claim(TOTAL_TARIF="1.2.3", LOS="-1")),
        pd.DataFrame(),
    )
    invalid = result.invalid_numeric_df
    overall = result.tariff_comparison_df.query("`Jenis Rawat` == 'Keseluruhan'").iloc[0]

    assert set(invalid.loc[invalid["Masalah"] == "Format tidak valid", "Field"]) >= {"TOTAL_TARIF", "LOS"}
    assert overall["Total Tarif INA-CBG"] is None
    assert overall["Klaim Tarif INA-CBG"] == 0


def test_export_treats_untrusted_text_as_text_not_formula() -> None:
    frame = _parsed(_claim(
        DPJP="=1+1",
        C2=json.dumps({"idrg": {"total_tarif": 1100, "cost_weight": 0.5, "grouper_version": "=1+1"}}),
    ))
    workbook = load_workbook(BytesIO(export_eklaim_analysis_to_excel(build_eklaim_analysis(frame, pd.DataFrame()))), data_only=False)

    dpjp_sheet = workbook["selisih_dpjp_ri"]
    dpjp_column = next(cell.column for cell in dpjp_sheet[1] if cell.value == "DPJP")
    assert dpjp_sheet.cell(2, dpjp_column).value == "=1+1"
    assert dpjp_sheet.cell(2, dpjp_column).data_type == "s"

    cmi_sheet = workbook["casemix_index"]
    version_row = next(row for row in range(2, cmi_sheet.max_row + 1)
                       if cmi_sheet.cell(row, 2).value == "Versi Bobot iDRG")
    assert cmi_sheet.cell(version_row, 3).value.startswith("=1+1")
    assert cmi_sheet.cell(version_row, 3).data_type == "s"


def test_export_keeps_tariff_amounts_numeric() -> None:
    result = build_eklaim_analysis(
        _parsed(_claim(TOTAL_TARIF="1000", C2=json.dumps({"idrg": {"total_tarif": 1100, "cost_weight": 0.5}}))),
        pd.DataFrame(),
    )
    workbook = load_workbook(BytesIO(export_eklaim_analysis_to_excel(result)), data_only=True)
    sheet = workbook["rekonsiliasi_tarif"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    overall_row = next(row for row in range(2, sheet.max_row + 1) if sheet.cell(row, headers["Jenis Rawat"]).value == "Keseluruhan")

    assert isinstance(sheet.cell(overall_row, headers["Total Tarif RS"]).value, (int, float))
    assert isinstance(sheet.cell(overall_row, headers["Total Tarif INA-CBG"]).value, (int, float))
    assert isinstance(sheet.cell(overall_row, headers["Total Tarif iDRG"]).value, (int, float))
